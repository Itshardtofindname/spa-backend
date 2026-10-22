
import io
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
from openpyxl import load_workbook

FILE_ID = "1SItN5xNdqOPipCp9QFHzwsKbE-G7NTTd"
CREDENTIALS_FILE = "google-service-account.json"
TARGET_SHEET = "Produktifitas"

credentials = service_account.Credentials.from_service_account_file(
    CREDENTIALS_FILE,
    scopes=["https://www.googleapis.com/auth/drive.readonly"],
)

drive = build(
    "drive",
    "v3",
    credentials=credentials,
    cache_discovery=False,
)

request = drive.files().get_media(fileId=FILE_ID)
buffer = io.BytesIO()
downloader = MediaIoBaseDownload(buffer, request)

done = False
while not done:
    status, done = downloader.next_chunk()
    if status:
        print(f"Download: {int(status.progress() * 100)}%")

buffer.seek(0)
workbook = load_workbook(
    buffer,
    read_only=True,
    data_only=True,
)

try:
    print("Sheets:", workbook.sheetnames)

    if TARGET_SHEET not in workbook.sheetnames:
        raise ValueError(f"Sheet {TARGET_SHEET!r} tidak ditemukan")

    sheet = workbook[TARGET_SHEET]

    print("Jumlah baris:", sheet.max_row)
    print("Jumlah kolom:", sheet.max_column)

    for row_number, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 10),
                        values_only=True),
        start=1,
    ):
        print(f"Baris {row_number}: {row[:8]}")

    print("\nVALIDASI SELESAI — tidak ada penulisan ke Supabase.")

finally:
    workbook.close()
