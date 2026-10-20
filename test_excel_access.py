
from pathlib import Path
from openpyxl import load_workbook

file_path = Path(
    r"C:\Users\Giovan\Downloads\RAPOT Produktifitas S003 - September 2026.xlsx"
)

print("File ditemukan:", file_path.exists())

try:
    with open(file_path, "rb") as file:
        print("File berhasil dibuka.")

    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True
    )

    print("Excel berhasil dibaca.")
    print("Daftar sheet:", workbook.sheetnames)
    workbook.close()

except Exception as error:
    print("Gagal:", repr(error))
