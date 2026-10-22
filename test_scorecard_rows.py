
import io
from collections import Counter

from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
from openpyxl import load_workbook

from services.scorecard_service import (
    find_header_row,
    normalize_header,
    normalize_npk,
    is_valid_salesman_row,
)

FILE_ID = "1SItN5xNdqOPipCp9QFHzwsKbE-G7NTTd"

credentials = service_account.Credentials.from_service_account_file(
    "google-service-account.json",
    scopes=["https://www.googleapis.com/auth/drive.readonly"],
)

drive = build("drive", "v3", credentials=credentials,
              cache_discovery=False)

request = drive.files().get_media(fileId=FILE_ID)
buffer = io.BytesIO()
downloader = MediaIoBaseDownload(buffer, request)

done = False
while not done:
    status, done = downloader.next_chunk()

buffer.seek(0)
workbook = load_workbook(buffer, read_only=True, data_only=True)

try:
    sheet = workbook["Produktifitas"]
    header_row = find_header_row(sheet)

    if header_row is None:
        raise ValueError("Header NPK, SALESMAN, PH tidak ditemukan")

    headers = [
        normalize_header(cell.value)
        for cell in sheet[header_row]
    ]

    npk_index = headers.index("npk")
    salesman_index = next(
        i for i, value in enumerate(headers)
        if value in {"salesman", "nama salesman"}
    )
    ph_index = headers.index("ph")

    valid_npks = []
    invalid_rows = 0

    for row in sheet.iter_rows(
        min_row=header_row + 1,
        values_only=True,
    ):
        normalized = [normalize_header(v) for v in row]

        if (
            "npk" in normalized
            and any(v in {"salesman", "nama salesman"} for v in normalized)
            and "ph" in normalized
        ):
            break

        npk = normalize_npk(
            row[npk_index] if npk_index < len(row) else None
        )
        name = (
            str(row[salesman_index]).strip()
            if salesman_index < len(row)
            and row[salesman_index] is not None
            else ""
        )
        ph = (
            str(row[ph_index]).strip().upper()
            if ph_index < len(row) and row[ph_index] is not None
            else ""
        )

        if is_valid_salesman_row(npk, name, ph):
            valid_npks.append(npk)
        else:
            invalid_rows += 1
            print(
                f"Baris dilewati: "
                f"NPK={npk!r}, "
                f"Nama={name!r}, "
                f"PH={ph!r}"
            )

    counts = Counter(valid_npks)
    duplicates = sum(count - 1 for count in counts.values() if count > 1)

    print("Header ditemukan pada baris:", header_row)
    print("Baris salesman valid:", len(valid_npks))
    print("NPK unik:", len(counts))
    print("Baris NPK duplikat:", duplicates)
    print("Baris dilewati:", invalid_rows)
    print("VALIDASI SELESAI — tidak ada penulisan ke Supabase.")

finally:
    workbook.close()
