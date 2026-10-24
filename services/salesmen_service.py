
from datetime import datetime, timezone
from io import BytesIO
import re

from openpyxl import load_workbook

from database import supabase


SUMM_SHEET = "SUMM"


def normalize_header(value):
    """Menyeragamkan teks header Excel."""
    if value is None:
        return ""

    text = str(value).strip().casefold()
    text = re.sub(r"[_./-]+", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def normalize_npk(value):
    """Mengubah NPK menjadi string konsisten."""
    if value is None or isinstance(value, bool):
        return ""

    if isinstance(value, float) and value.is_integer():
        value = int(value)

    text = str(value).strip()

    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".")[0]

    return text


def parse_summ_number(value):
    """Mengubah nilai indikator SUMM menjadi angka."""
    if value is None or isinstance(value, bool):
        return None

    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip().replace(" ", "")

    if not text:
        return None

    if "," in text and "." not in text:
        text = text.replace(",", ".")

    try:
        return float(text)
    except (ValueError, TypeError):
        return None


def _find_summ_headers(sheet):
    """
    Mencari dua baris header pada sheet SUMM.

    Baris pertama dapat berisi kelompok indikator seperti RESULT,
    STRATEGIC ALLIGNMENT, PROCESS, dan FINAL.

    Baris berikutnya berisi nama kolom sebenarnya, seperti NPK,
    NAMA, SALES ACHIEV & GR, CROSS - UP SELL, GRADE, dan SCORE.
    """
    max_header_row = min(sheet.max_row or 0, 40)

    for row_number in range(1, max_header_row + 1):
        current = [
            normalize_header(cell.value)
            for cell in sheet[row_number]
        ]

        next_row = (
            [
                normalize_header(cell.value)
                for cell in sheet[row_number + 1]
            ]
            if row_number < sheet.max_row
            else []
        )

        # Periksa kedua baris, bukan hanya header yang sudah digabung.
        all_headers = set(current + next_row)

        has_npk = "npk" in all_headers
        has_name = any(
            value in {
                "nama",
                "name",
                "salesman",
                "nama salesman",
            }
            for value in all_headers
        )
        has_grade = "grade" in all_headers
        has_score = "score" in all_headers
        has_result = any(
            "sales achiev" in value
            or "sales achievement" in value
            for value in all_headers
        )

        if not (
            has_npk
            and has_name
            and has_grade
            and has_score
            and has_result
        ):
            continue

        columns = {}

        max_columns = max(len(current), len(next_row))

        for index in range(max_columns):
            first_value = (
                current[index]
                if index < len(current)
                else ""
            )
            second_value = (
                next_row[index]
                if index < len(next_row)
                else ""
            )

            # Prioritaskan baris kedua karena berisi nama indikator
            # yang sebenarnya. Gunakan baris pertama sebagai fallback.
            selected_header = second_value or first_value

            if selected_header:
                columns[index] = selected_header

        return row_number, columns

    raise ValueError(
        "Header SUMM tidak ditemukan. "
        "Pastikan sheet memiliki NPK, NAMA, SALES ACHIEV, "
        "GRADE, dan SCORE."
    )


def _get_summ_column(columns, aliases):
    """Mencari kolom menggunakan daftar alias."""
    for index, header in columns.items():
        for alias in aliases:
            if alias in header:
                return index

    return None


def _get_cell(row, index):
    """Mengambil nilai kolom dengan aman."""
    if index is None or index >= len(row):
        return None

    return row[index]


def _extract_summ_records(
    file_bytes,
    sales_office,
    report_period,
    source_file,
):
    """Membaca sheet SUMM dan menghasilkan record terstandardisasi."""
    workbook = load_workbook(
        BytesIO(file_bytes),
        read_only=True,
        data_only=True,
    )

    try:
        if SUMM_SHEET not in workbook.sheetnames:
            raise ValueError(
                f"Sheet '{SUMM_SHEET}' tidak ditemukan. "
                f"Sheet tersedia: {', '.join(workbook.sheetnames)}"
            )

        sheet = workbook[SUMM_SHEET]
        header_row, columns = _find_summ_headers(sheet)

        column_map = {
            "npk": _get_summ_column(
                columns,
                ["npk"],
            ),
            "nama": _get_summ_column(
                columns,
                ["nama", "salesman"],
            ),
            "result_sales_achiev_gr": _get_summ_column(
                columns,
                ["sales achiev", "sales achievement"],
            ),
            "cross_up_sell": _get_summ_column(
                columns,
                ["cross up sell", "cross selling", "up sell"],
            ),
            "active_outlet": _get_summ_column(
                columns,
                ["active outlet"],
            ),
            "repeat_order": _get_summ_column(
                columns,
                ["repeat order"],
            ),
            "process_dks": _get_summ_column(
                columns,
                ["dks"],
            ),
            "final_grade": _get_summ_column(
                columns,
                ["grade"],
            ),
            "final_score": _get_summ_column(
                columns,
                ["score"],
            ),
        }

        required_columns = [
            "npk",
            "nama",
            "result_sales_achiev_gr",
            "cross_up_sell",
            "active_outlet",
            "repeat_order",
            "process_dks",
            "final_grade",
            "final_score",
        ]

        missing = [
            key
            for key in required_columns
            if column_map[key] is None
        ]

        if missing:
            raise ValueError(
                "Kolom SUMM tidak lengkap: "
                + ", ".join(missing)
            )

        records_by_npk = {}
        duplicate_npks = set()
        skipped_rows = 0

        # Pada workbook saat ini, baris pertama data berada
        # dua baris setelah baris kelompok header.
        for row in sheet.iter_rows(
            min_row=header_row + 2,
            values_only=True,
        ):
            npk = normalize_npk(
                _get_cell(row, column_map["npk"])
            )

            name_value = _get_cell(
                row,
                column_map["nama"],
            )

            nama = (
                str(name_value).strip()
                if name_value is not None
                else ""
            )

            if not npk or not nama:
                skipped_rows += 1
                continue

            if not re.fullmatch(r"\d+", npk):
                skipped_rows += 1
                continue

            result_value = parse_summ_number(
                _get_cell(
                    row,
                    column_map["result_sales_achiev_gr"],
                )
            )

            cross_value = parse_summ_number(
                _get_cell(
                    row,
                    column_map["cross_up_sell"],
                )
            )

            outlet_value = parse_summ_number(
                _get_cell(
                    row,
                    column_map["active_outlet"],
                )
            )

            repeat_value = parse_summ_number(
                _get_cell(
                    row,
                    column_map["repeat_order"],
                )
            )

            dks_value = parse_summ_number(
                _get_cell(
                    row,
                    column_map["process_dks"],
                )
            )

            grade_value = parse_summ_number(
                _get_cell(
                    row,
                    column_map["final_grade"],
                )
            )

            score_value = _get_cell(
                row,
                column_map["final_score"],
            )

            final_score = (
                str(score_value).strip().upper()
                if score_value is not None
                else ""
            )

            metric_values = [
                result_value,
                cross_value,
                outlet_value,
                repeat_value,
                dks_value,
                grade_value,
            ]

            # Abaikan baris yang memang belum mempunyai nilai SUMM.
            if (
                all(value is None for value in metric_values)
                and not final_score
            ):
                skipped_rows += 1
                continue

            # Jangan menyimpan nilai parsial sebagai data lengkap.
            if grade_value is None or not final_score:
                raise ValueError(
                    f"Grade atau score SUMM kosong untuk "
                    f"NPK {npk} ({nama})."
                )

            record = {
                "npk": npk,
                "nama": nama,
                "sales_office": sales_office,
                "report_period": str(report_period),
                "final_grade": (
                    int(grade_value)
                    if grade_value.is_integer()
                    else grade_value
                ),
                "final_score": final_score,
                "strategic_alignment": {
                    "cross_up_sell": cross_value or 0,
                    "active_outlet": outlet_value or 0,
                    "repeat_order": repeat_value or 0,
                    "process_dks": dks_value or 0,
                },
                "result_sales_achiev_gr": result_value or 0,
                "source_file": source_file,
                "updated_at": datetime.now(
                    timezone.utc
                ).isoformat(),
            }

            if npk in records_by_npk:
                duplicate_npks.add(npk)
                continue

            records_by_npk[npk] = record

        if duplicate_npks:
            raise ValueError(
                "Impor SUMM dibatalkan karena NPK duplikat: "
                + ", ".join(sorted(duplicate_npks))
            )

        records = list(records_by_npk.values())

        if not records:
            raise ValueError(
                "Tidak ada baris SUMM valid yang ditemukan."
            )

        return records, skipped_rows

    finally:
        workbook.close()


def import_salesman_summ_excel(
    file_bytes,
    sales_office,
    report_period,
    source_file,
    dry_run=False,
    db_client=None,
):
    """
    Impor sheet SUMM ke salesmen.summ_history.

    Key histori:
        YYYY-MM-DD|SALES OFFICE

    Histori periode lain dipertahankan.
    Data master salesman yang sudah ada tidak ditimpa,
    kecuali kolom summ_history.
    """
    sales_office = str(sales_office or "").strip()
    report_period = str(report_period or "").strip()

    if not sales_office:
        raise ValueError("sales_office tidak boleh kosong.")

    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", report_period):
        raise ValueError(
            "report_period harus berformat YYYY-MM-DD."
        )

    try:
        datetime.strptime(report_period, "%Y-%m-%d")
    except ValueError as exc:
        raise ValueError(
            "report_period bukan tanggal yang valid."
        ) from exc

    records, skipped_rows = _extract_summ_records(
        file_bytes=file_bytes,
        sales_office=sales_office,
        report_period=report_period,
        source_file=source_file,
    )

    if dry_run:
        return {
            "message": (
                "Validasi SUMM berhasil; database tidak diubah."
            ),
            "dry_run": True,
            "sheet": SUMM_SHEET,
            "summ_rows_imported": len(records),
            "skipped_rows": skipped_rows,
            "sales_office": sales_office,
            "report_period": report_period,
            "source_file": source_file,
            "preview": records[:5],
        }

    client = db_client if db_client is not None else supabase
    history_key = f"{report_period}|{sales_office}"

    inserted = 0
    updated = 0

    for record in records:
        existing_rows = (
            client.table("salesmen")
            .select(
                "npk,nama,sales_office,grup_product,summ_history"
            )
            .eq("npk", record["npk"])
            .limit(1)
            .execute()
            .data
            or []
        )

        history_entry = {
            "npk": record["npk"],
            "nama": record["nama"],
            "sales_office": sales_office,
            "report_period": report_period,
            "final_grade": record["final_grade"],
            "final_score": record["final_score"],
            "strategic_alignment": (
                record["strategic_alignment"]
            ),
            "result_sales_achiev_gr": (
                record["result_sales_achiev_gr"]
            ),
            "source_file": source_file,
            "updated_at": record["updated_at"],
        }

        if existing_rows:
            existing = existing_rows[0]
            history = existing.get("summ_history") or {}

            if not isinstance(history, dict):
                raise ValueError(
                    f"summ_history untuk NPK {record['npk']} "
                    "bukan objek JSON. Periksa datanya sebelum impor."
                )

            # Salin histori agar periode lain tidak hilang.
            new_history = dict(history)
            new_history[history_key] = history_entry

            (
                client.table("salesmen")
                .update({
                    "summ_history": new_history,
                })
                .eq("npk", record["npk"])
                .execute()
            )

            updated += 1

        else:
            # Jika NPK belum ada, buat master salesman.
            # Pastikan kolom lain pada skema database mengizinkan
            # nilai default atau NULL untuk field yang tidak disertakan.
            (
                client.table("salesmen")
                .insert({
                    "npk": record["npk"],
                    "nama": record["nama"],
                    "sales_office": sales_office,
                    "summ_history": {
                        history_key: history_entry,
                    },
                })
                .execute()
            )

            inserted += 1

    return {
        "message": "Impor SUMM berhasil.",
        "dry_run": False,
        "sheet": SUMM_SHEET,
        "summ_rows_imported": len(records),
        "summ_rows_inserted": inserted,
        "summ_rows_updated": updated,
        "skipped_rows": skipped_rows,
        "sales_office": sales_office,
        "report_period": report_period,
        "source_file": source_file,
    }


def get_all_salesmen():
    response = (
        supabase.table("salesmen")
        .select("*")
        .execute()
    )
    return response.data or []


def get_filter_options():
    response = (
        supabase.table("salesmen")
        .select("sales_office, grup_product, nama, npk")
        .execute()
    )

    raw_data = response.data or []

    return {
        "raw": raw_data,
        "offices": sorted({
            item["sales_office"]
            for item in raw_data
            if item.get("sales_office")
        }),
        "groups": sorted({
            item["grup_product"]
            for item in raw_data
            if item.get("grup_product")
        }),
        "salesmen": sorted({
            item["nama"]
            for item in raw_data
            if item.get("nama")
        }),
    }


def get_salesman_performance(
    office: str | None = None,
    group: str | None = None,
    salesman: str | None = None,
):
    query = supabase.table("salesmen").select("*, targets(*)")

    if salesman and salesman.strip():
        query = query.ilike(
            "nama",
            f"%{salesman.strip()}%",
        )

    if office and office.strip():
        query = query.eq(
            "sales_office",
            office.strip(),
        )

    if group and group.strip():
        query = query.eq(
            "grup_product",
            group.strip(),
        )

    response = query.execute()
    return response.data or []


def get_salesman_summ(
    office: str | None = None,
    group: str | None = None,
    salesman: str | None = None,
    report_period: str | None = None,
):
    """Membaca histori SUMM dari salesmen.summ_history."""
    query = supabase.table("salesmen").select(
        "npk, nama, sales_office, grup_product, summ_history"
    )

    if salesman and salesman.strip():
        query = query.ilike(
            "nama",
            f"%{salesman.strip()}%",
        )

    if office and office.strip():
        query = query.eq(
            "sales_office",
            office.strip(),
        )

    if group and group.strip():
        query = query.eq(
            "grup_product",
            group.strip(),
        )

    response = query.execute()
    rows = response.data or []

    results = []

    for row in rows:
        history = row.get("summ_history") or {}

        if not isinstance(history, dict):
            continue

        for entry in history.values():
            if not isinstance(entry, dict):
                continue

            if (
                report_period
                and str(
                    entry.get("report_period", "")
                ).strip()
                != str(report_period).strip()
            ):
                continue

            if office and (
                str(
                    entry.get("sales_office", "")
                ).strip()
                != office.strip()
            ):
                continue

            results.append({
                **entry,
                "npk": entry.get("npk") or row.get("npk"),
                "nama": entry.get("nama") or row.get("nama"),
                "sales_office": (
                    entry.get("sales_office")
                    or row.get("sales_office")
                ),
                "grup_product": row.get("grup_product"),
            })

    results.sort(
        key=lambda item: (
            str(item.get("nama", "")).casefold(),
            str(item.get("npk", "")),
        )
    )

    return results
