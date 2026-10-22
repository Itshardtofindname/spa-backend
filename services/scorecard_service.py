import calendar
import re
from datetime import date, datetime, timezone
from io import BytesIO

from openpyxl import load_workbook

from database import supabase


TARGET_SHEET = "Produktifitas"


def normalize_header(value):
    """Menyeragamkan nama kolom Excel."""
    if value is None:
        return ""

    text = str(value).strip().lower()
    text = re.sub(r"[_./-]+", " ", text)
    text = re.sub(r"\s+", " ", text)

    return text.strip()


def normalize_npk(value):
    """Mengubah NPK menjadi string yang konsisten."""
    if value is None or isinstance(value, bool):
        return ""

    if isinstance(value, float) and value.is_integer():
        value = int(value)

    text = str(value).strip()

    # Menangani NPK yang terbaca sebagai 1584.0.
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".")[0]

    return text


def normalize_name(value):
    """Menyeragamkan nama untuk pencocokan data lama."""
    if value is None:
        return ""

    return re.sub(
        r"\s+",
        " ",
        str(value).strip().casefold(),
    )


def parse_number(value):
    """Mengubah nilai Excel menjadi float."""
    if value is None or isinstance(value, bool):
        return None

    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip().replace(" ", "")

    if not text:
        return None

    if "," in text and "." not in text:
        text = text.replace(",", ".")

    try:
        return float(text)
    except (ValueError, TypeError):
        return None


def get_month_label(value):
    """Mengubah header bulan menjadi format Jan-26."""
    if isinstance(value, (datetime, date)):
        return value.strftime("%b-%y")

    if value is None:
        return None

    text = re.sub(r"\s+", " ", str(value).strip())

    match = re.match(
        r"^([A-Za-z]+)[- /]+(\d{2,4})$",
        text,
    )

    if not match:
        return None

    month_text = match.group(1).lower()
    year_text = match.group(2)

    month_number = next(
        (
            number
            for number in range(1, 13)
            if month_text
            in {
                calendar.month_abbr[number].lower(),
                calendar.month_name[number].lower(),
            }
        ),
        None,
    )

    if month_number is None:
        return None

    year_number = int(year_text)

    if len(year_text) == 2:
        year_number += 2000

    try:
        return date(
            year_number,
            month_number,
            1,
        ).strftime("%b-%y")
    except ValueError:
        return None


def find_header_row(sheet):
    """Mencari baris header NPK, SALESMAN, dan PH."""
    for row_number in range(1, min(sheet.max_row, 30) + 1):
        values = [
            normalize_header(cell.value)
            for cell in sheet[row_number]
        ]

        has_npk = "npk" in values
        has_salesman = any(
            value in {"salesman", "nama salesman"}
            for value in values
        )
        has_ph = "ph" in values

        if has_npk and has_salesman and has_ph:
            return row_number

    return None


def is_valid_salesman_row(npk, name, ph):
    """Memastikan baris merupakan data salesman."""
    if not npk or not name or not ph:
        return False

    if not re.fullmatch(r"\d+", npk):
        return False

    if normalize_header(name) in {
        "total",
        "subtotal",
        "grand total",
        "high",
        "medium",
        "low",
    }:
        return False

    if not re.fullmatch(r"TL\d+[A-Z]?", ph.upper()):
        return False

    return True


def get_existing_records(
    sales_office,
    report_period,
    db_client=None,
):
    """Mengambil identitas data lama untuk kantor dan periode yang sama."""
    client = db_client if db_client is not None else supabase

    response = (
        client.table("scorecard_productivity")
        .select("id, npk, salesman")
        .eq("sales_office", sales_office)
        .eq("report_period", report_period)
        .execute()
    )

    return response.data or []


def import_scorecard_excel(
    file_bytes,
    sales_office,
    report_period,
    source_file,
    dry_run=False,
    db_client=None,
):
    """
    Import data Excel dengan memperbarui row yang sudah ada.

    Strategi pencocokan:
    1. Cocokkan NPK terlebih dahulu.
    2. Jika NPK berubah, cocokkan nama yang unik.
    3. Pertahankan UUID lama jika pasangan ditemukan.
    4. Buat row baru jika tidak ada pasangan.
    5. Tidak menghapus row lama yang tidak ditemukan di Excel.

    Mode dry_run tidak membaca atau menulis database.
    """
    workbook = load_workbook(
        BytesIO(file_bytes),
        read_only=True,
        data_only=True,
    )

    try:
        if TARGET_SHEET not in workbook.sheetnames:
            raise ValueError(
                f"Sheet '{TARGET_SHEET}' tidak ditemukan. "
                f"Sheet tersedia: {', '.join(workbook.sheetnames)}"
            )

        sheet = workbook[TARGET_SHEET]
        header_row = find_header_row(sheet)

        if header_row is None:
            raise ValueError(
                "Header tidak ditemukan. Pastikan sheet memiliki "
                "kolom NPK, SALESMAN, dan PH."
            )

        header_cells = [
            cell.value for cell in sheet[header_row]
        ]

        headers = [
            normalize_header(value)
            for value in header_cells
        ]

        columns = {}

        for index, header in enumerate(headers):
            if header and header not in columns:
                columns[header] = index

        month_columns = {}

        for index, value in enumerate(header_cells):
            month_label = get_month_label(value)

            if month_label and month_label not in month_columns:
                month_columns[month_label] = index

        def get_value(row, *aliases):
            for alias in aliases:
                index = columns.get(normalize_header(alias))

                if index is not None and index < len(row):
                    return row[index]

            return None

        records_by_npk = {}
        duplicate_npks = set()
        skipped_rows = 0

        for row in sheet.iter_rows(
            min_row=header_row + 1,
            values_only=True,
        ):
            normalized_row = [
                normalize_header(value)
                for value in row
            ]

            # Berhenti jika tabel berikutnya dimulai.
            if (
                "npk" in normalized_row
                and any(
                    value in {"salesman", "nama salesman"}
                    for value in normalized_row
                )
                and "ph" in normalized_row
            ):
                break

            npk = normalize_npk(
                get_value(row, "npk")
            )

            name_value = get_value(
                row,
                "salesman",
                "nama salesman",
            )

            name = (
                str(name_value).strip()
                if name_value is not None
                else ""
            )

            ph_value = get_value(row, "ph")

            ph = (
                str(ph_value).strip().upper()
                if ph_value is not None
                else ""
            )

            if not is_valid_salesman_row(npk, name, ph):
                skipped_rows += 1
                continue

            monthly_values = {}

            for month, index in month_columns.items():
                if index >= len(row):
                    continue

                value = parse_number(row[index])

                if value is not None:
                    monthly_values[month] = value

            record = {
                "sales_office": sales_office,
                "report_period": report_period,
                "npk": npk,
                "salesman": name,
                "ph": ph,
                "monthly_values": monthly_values,
                "q1": parse_number(
                    get_value(row, "q1")
                ),
                "perform": parse_number(
                    get_value(row, "perform")
                ),
                "oa": parse_number(
                    get_value(row, "oa")
                ),
                "r_order": parse_number(
                    get_value(row, "r order", "rorder")
                ),
                "cross_sell": parse_number(
                    get_value(row, "cross sell", "crosssell")
                ),
                "dks": parse_number(
                    get_value(row, "dks")
                ),
                "source_file": source_file,
                "updated_at": datetime.now(
                    timezone.utc
                ).isoformat(),
            }

            if npk in records_by_npk:
                duplicate_npks.add(npk)
                continue

            records_by_npk[npk] = record

        if duplicate_npks:
            raise ValueError(
                "Impor dibatalkan karena ditemukan NPK duplikat: "
                + ", ".join(sorted(duplicate_npks))
            )

        records = list(records_by_npk.values())

        if not records:
            raise ValueError(
                "Tidak ada baris salesman valid yang ditemukan. "
                "Periksa header dan format kolom Excel."
            )

        # Mode validasi: tidak membaca atau menulis database.
        if dry_run:
            return {
                "message": (
                    "Validasi berhasil; database tidak diakses"
                ),
                "dry_run": True,
                "sheet": TARGET_SHEET,
                "sheet_rows_imported": len(records),
                "skipped_rows": skipped_rows,
                "sales_office": sales_office,
                "report_period": str(report_period),
                "source_file": source_file,
                "preview": [
                    {
                        "npk": record["npk"],
                        "salesman": record["salesman"],
                        "ph": record["ph"],
                        "monthly_values": record[
                            "monthly_values"
                        ],
                    }
                    for record in records[:5]
                ],
            }

        # Client dipilih setelah dry-run agar validasi tidak mengakses DB.
        client = db_client if db_client is not None else supabase

        existing_records = get_existing_records(
            sales_office,
            report_period,
            db_client=client,
        )

        existing_by_npk = {}
        existing_by_name = {}

        for existing in existing_records:
            old_npk = normalize_npk(
                existing.get("npk")
            )
            old_name = normalize_name(
                existing.get("salesman")
            )

            if old_npk:
                existing_by_npk.setdefault(
                    old_npk, []
                ).append(existing)

            if old_name:
                existing_by_name.setdefault(
                    old_name, []
                ).append(existing)

        matched_existing_ids = set()
        prepared_records = []
        matched_by_npk = 0
        matched_by_name = 0
        new_records = 0

        for record in records:
            npk = record["npk"]
            name_key = normalize_name(
                record["salesman"]
            )

            # Prioritas pertama: NPK sama.
            npk_matches = existing_by_npk.get(npk, [])

            if len(npk_matches) > 1:
                raise ValueError(
                    f"Ditemukan beberapa row lama dengan NPK {npk}. "
                    "Bersihkan data duplikat sebelum import."
                )

            existing = (
                npk_matches[0]
                if npk_matches
                else None
            )

            if existing is not None:
                if existing["id"] in matched_existing_ids:
                    raise ValueError(
                        f"Row lama untuk NPK {npk} cocok dengan "
                        "lebih dari satu baris Excel."
                    )

                matched_by_npk += 1

            else:
                # Prioritas kedua: nama unik jika NPK berubah.
                name_matches = existing_by_name.get(
                    name_key, []
                )

                name_matches = [
                    item
                    for item in name_matches
                    if item["id"] not in matched_existing_ids
                ]

                if len(name_matches) > 1:
                    raise ValueError(
                        f"Nama '{record['salesman']}' cocok dengan "
                        "beberapa row lama. Import dibatalkan untuk "
                        "menghindari perubahan pada row yang salah."
                    )

                if len(name_matches) == 1:
                    existing = name_matches[0]
                    matched_by_name += 1

            if existing is not None:
                # Cegah NPK baru mengambil identitas row lain.
                conflicting_npk = existing_by_npk.get(
                    npk, []
                )

                if any(
                    item["id"] != existing["id"]
                    for item in conflicting_npk
                ):
                    raise ValueError(
                        f"NPK {npk} sudah dimiliki row lain pada "
                        "kantor dan periode ini. Import dibatalkan "
                        "untuk mencegah konflik data."
                    )

                record["id"] = existing["id"]
                matched_existing_ids.add(existing["id"])

            else:
                new_records += 1

            prepared_records.append(record)

        # Upsert menggunakan client yang sama dengan proses pencocokan.
        (
            client.table("scorecard_productivity")
            .upsert(
                prepared_records,
                on_conflict="id",
            )
            .execute()
        )

        return {
            "message": "Impor Scorecard berhasil",
            "sheet": TARGET_SHEET,
            "sheet_rows_imported": len(prepared_records),
            "updated_by_npk": matched_by_npk,
            "updated_by_name": matched_by_name,
            "new_rows": new_records,
            "skipped_rows": skipped_rows,
            "sales_office": sales_office,
            "report_period": str(report_period),
            "source_file": source_file,
        }

    finally:
        workbook.close()


def get_scorecard_productivity(
    sales_office,
    report_period,
):
    """Mengambil Scorecard berdasarkan kantor dan periode."""
    response = (
        supabase.table("scorecard_productivity")
        .select("*")
        .eq("sales_office", sales_office)
        .eq("report_period", report_period)
        .order("ph")
        .order("npk")
        .execute()
    )

    return response.data or []