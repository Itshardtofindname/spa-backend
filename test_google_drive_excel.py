
import io
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
from openpyxl import load_workbook

FILE_ID = "1SItN5xNdqOPipCp9QFHzwsKbE-G7NTTd"
CREDENTIALS_FILE = "google-service-account.json"

credentials = service_account.Credentials.from_service_account_file(
    CREDENTIALS_FILE,
    scopes=["https://www.googleapis.com/auth/drive.readonly"],
)

drive = build(
    "drive",
    "v3",
    credentials=credentials,
    cache_discovery=False,
)

request = drive.files().get_media(fileId=FILE_ID)
excel_buffer = io.BytesIO()
downloader = MediaIoBaseDownload(excel_buffer, request)

done = False
while not done:
    status, done = downloader.next_chunk()
    if status:
        print(f"Download: {int(status.progress() * 100)}%")

excel_buffer.seek(0)
workbook = load_workbook(
    excel_buffer,
    read_only=True,
    data_only=True,
)

print("\nNama sheet:", workbook.sheetnames)

for sheet in workbook.worksheets:
    print(
        f"Sheet: {sheet.title} | "
        f"Baris: {sheet.max_row} | "
        f"Kolom: {sheet.max_column}"
    )

if "Produktifitas" in workbook.sheetnames:
    print("\nBERHASIL: Sheet Produktifitas ditemukan.")
else:
    print("\nPERHATIAN: Sheet Produktifitas tidak ditemukan.")

workbook.close()
